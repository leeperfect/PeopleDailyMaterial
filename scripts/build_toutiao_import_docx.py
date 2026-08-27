#!/usr/bin/env python3
"""将项目中的成稿 Markdown 转成适合今日头条“文档导入”的 DOCX。"""

from __future__ import annotations

import argparse
import json
import re
import tempfile
from pathlib import Path

from PIL import Image
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


IMAGE_RE = re.compile(r"^!\[(?P<alt>[^]]*)\]\((?P<path>[^)]+)\)\s*$")
HEADING_RE = re.compile(r"^(?P<level>#{1,3})\s+(?P<text>.+?)\s*$")
ORDERED_RE = re.compile(r"^\s*\d+[.)]\s+(?P<text>.+?)\s*$")
BULLET_RE = re.compile(r"^\s*[-*+]\s+(?P<text>.+?)\s*$")
BOLD_RE = re.compile(r"\*\*(.+?)\*\*")


def set_run_font(run, name: str, size: float | None = None, bold: bool | None = None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def add_inline_markdown(paragraph, text: str, *, font_name: str = "Hiragino Sans GB", size: float = 12):
    text = text.replace(r"\*\*", "**").replace(r"\*", "*")
    cursor = 0
    for match in BOLD_RE.finditer(text):
        if match.start() > cursor:
            run = paragraph.add_run(text[cursor : match.start()])
            set_run_font(run, font_name, size)
        run = paragraph.add_run(match.group(1))
        set_run_font(run, font_name, size, True)
        cursor = match.end()
    if cursor < len(text):
        run = paragraph.add_run(text[cursor:])
        set_run_font(run, font_name, size)


def add_numbering_definition(document: Document, fmt: str, text: str) -> int:
    numbering = document.part.numbering_part.element
    abstract_ids = [int(e.get(qn("w:abstractNumId"))) for e in numbering.findall(qn("w:abstractNum"))]
    abstract_id = max(abstract_ids, default=-1) + 1
    num_ids = [int(e.get(qn("w:numId"))) for e in numbering.findall(qn("w:num"))]
    num_id = max(num_ids, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)

    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    level.append(start)
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), fmt)
    level.append(num_fmt)
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), text)
    level.append(lvl_text)
    justification = OxmlElement("w:lvlJc")
    justification.set(qn("w:val"), "left")
    level.append(justification)
    ppr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "540")
    tabs.append(tab)
    ppr.append(tabs)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "540")
    ind.set(qn("w:hanging"), "360")
    ppr.append(ind)
    level.append(ppr)
    abstract.append(level)
    numbering.append(abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def apply_numbering(paragraph, num_id: int):
    ppr = paragraph._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num = OxmlElement("w:numId")
    num.set(qn("w:val"), str(num_id))
    num_pr.append(ilvl)
    num_pr.append(num)
    ppr.append(num_pr)


def set_paragraph_spacing(paragraph, *, before=0, after=8, line=1.45):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line


def prepare_image(source: Path, destination: Path):
    with Image.open(source) as image:
        converted = image.convert("RGB")
        converted.thumbnail((1600, 1600), Image.Resampling.LANCZOS)
        converted.save(destination, "JPEG", quality=84, optimize=True, progressive=True)


def configure_document(document: Document):
    section = document.sections[0]
    section.start_type = WD_SECTION.NEW_PAGE
    section.top_margin = Inches(0.78)
    section.bottom_margin = Inches(0.78)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Hiragino Sans GB"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Hiragino Sans GB")
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 1.45
    normal.paragraph_format.space_after = Pt(8)

    for style_name, size in (("Title", 20), ("Heading 1", 16), ("Heading 2", 14)):
        style = styles[style_name]
        style.font.name = "Hiragino Sans GB"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Hiragino Sans GB")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(24, 24, 24)

    if "Toutiao Quote" not in styles:
        quote = styles.add_style("Toutiao Quote", WD_STYLE_TYPE.PARAGRAPH)
        quote.font.name = "Hiragino Sans GB"
        quote._element.rPr.rFonts.set(qn("w:eastAsia"), "Hiragino Sans GB")
        quote.font.size = Pt(11.5)
        quote.font.color.rgb = RGBColor(70, 70, 70)
        quote.paragraph_format.left_indent = Inches(0.32)
        quote.paragraph_format.right_indent = Inches(0.15)
        quote.paragraph_format.space_before = Pt(4)
        quote.paragraph_format.space_after = Pt(10)
        quote.paragraph_format.line_spacing = 1.4


def strip_frontmatter(lines: list[str]) -> list[str]:
    if not lines or lines[0].strip() != "---":
        return lines
    for index in range(1, len(lines)):
        if lines[index].strip() == "---":
            return lines[index + 1 :]
    return lines


def build(source: Path, output: Path, project_root: Path) -> dict[str, int | str]:
    lines = strip_frontmatter(source.read_text(encoding="utf-8").splitlines())
    document = Document()
    configure_document(document)

    active_list_kind: str | None = None
    active_num_id: int | None = None
    in_code = False
    code_lines: list[str] = []
    image_count = 0
    ordered_item_count = 0
    bullet_item_count = 0
    heading_count = 0

    with tempfile.TemporaryDirectory(prefix="toutiao-docx-") as temp_dir:
        temp_root = Path(temp_dir)
        for raw_line in lines:
            line = raw_line.rstrip()
            if line.startswith("```"):
                if not in_code:
                    in_code = True
                    code_lines = []
                else:
                    num_id = add_numbering_definition(document, "decimal", "%1.")
                    for item in code_lines:
                        match = ORDERED_RE.match(item)
                        text = match.group("text") if match else item
                        paragraph = document.add_paragraph()
                        apply_numbering(paragraph, num_id)
                        set_paragraph_spacing(paragraph, after=5, line=1.35)
                        add_inline_markdown(paragraph, text)
                    in_code = False
                    code_lines = []
                active_list_kind = None
                active_num_id = None
                continue
            if in_code:
                if line.strip():
                    code_lines.append(line.strip())
                continue
            if not line.strip():
                continue

            image_match = IMAGE_RE.match(line.strip())
            if image_match:
                image_path = (project_root / image_match.group("path")).resolve()
                if not image_path.exists():
                    raise FileNotFoundError(f"找不到配图：{image_path}")
                compressed_path = temp_root / f"image-{len(list(temp_root.iterdir())) + 1:02d}.jpg"
                prepare_image(image_path, compressed_path)
                paragraph = document.add_paragraph()
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                paragraph.paragraph_format.space_before = Pt(4)
                paragraph.paragraph_format.space_after = Pt(10)
                paragraph.add_run().add_picture(str(compressed_path), width=Inches(6.45))
                image_count += 1
                active_list_kind = None
                active_num_id = None
                continue

            heading_match = HEADING_RE.match(line)
            if heading_match:
                level = len(heading_match.group("level"))
                text = heading_match.group("text")
                if level == 1:
                    # 今日头条的“文档导入”不会用 H1 自动填写标题栏，而会把它
                    # 作为正文第一段。标题由浏览器流程单独填写，因此这里跳过 H1，
                    # 避免导入后正文重复出现文章标题。
                    active_list_kind = None
                    active_num_id = None
                    continue
                else:
                    paragraph = document.add_paragraph(style="Heading 1" if level == 2 else "Heading 2")
                    paragraph.paragraph_format.space_before = Pt(16)
                    paragraph.paragraph_format.space_after = Pt(8)
                    add_inline_markdown(paragraph, text, font_name="Hiragino Sans GB", size=16 if level == 2 else 14)
                    heading_count += 1
                active_list_kind = None
                active_num_id = None
                continue

            if line.startswith(">"):
                paragraph = document.add_paragraph(style="Toutiao Quote")
                add_inline_markdown(paragraph, line[1:].strip(), size=11.5)
                active_list_kind = None
                active_num_id = None
                continue

            ordered_match = ORDERED_RE.match(line)
            bullet_match = BULLET_RE.match(line)
            if ordered_match or bullet_match:
                kind = "ordered" if ordered_match else "bullet"
                if active_list_kind != kind or active_num_id is None:
                    active_num_id = add_numbering_definition(
                        document,
                        "decimal" if kind == "ordered" else "bullet",
                        "%1." if kind == "ordered" else "•",
                    )
                    active_list_kind = kind
                paragraph = document.add_paragraph()
                apply_numbering(paragraph, active_num_id)
                set_paragraph_spacing(paragraph, after=5, line=1.35)
                add_inline_markdown(paragraph, (ordered_match or bullet_match).group("text"))
                if kind == "ordered":
                    ordered_item_count += 1
                else:
                    bullet_item_count += 1
                continue

            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            set_paragraph_spacing(paragraph)
            add_inline_markdown(paragraph, line.strip())
            active_list_kind = None
            active_num_id = None

        output.parent.mkdir(parents=True, exist_ok=True)
        document.save(output)
    size_bytes = output.stat().st_size
    if size_bytes > 15 * 1024 * 1024:
        raise RuntimeError(
            f"今日头条导入文档为 {size_bytes / 1024 / 1024:.2f} MB，超过 15 MB 限制"
        )
    return {
        "docx_path": str(output),
        "size_bytes": size_bytes,
        "image_count": image_count,
        "heading_count": heading_count,
        "ordered_item_count": ordered_item_count,
        "bullet_item_count": bullet_item_count,
    }


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    project_root = Path(__file__).resolve().parents[1]
    result = build(args.source.resolve(), args.output.resolve(), project_root)
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
