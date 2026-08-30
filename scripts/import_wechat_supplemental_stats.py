#!/usr/bin/env python3
"""导入微信公众号搜索与粉丝补充数据。"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
import shutil
import sqlite3
import subprocess
import sys
from datetime import datetime
from pathlib import Path
from urllib.parse import parse_qs, urlparse

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_DB_PATH = ROOT / "data" / "core" / "wechat_official_account.sqlite"
DATA_ROOT = ROOT / "data" / "data_analysis"


def now() -> str:
    return datetime.now().isoformat(timespec="seconds")


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def parse_int(value: object) -> int | None:
    text = str(value or "").strip().replace(",", "")
    if not text or text == "-":
        return None
    return int(float(text))


def parse_float(value: object) -> float | None:
    text = str(value or "").strip().replace(",", "")
    if not text or text == "-":
        return None
    if text.endswith("%"):
        return float(text[:-1]) / 100
    return float(text)


def read_csv(path: Path) -> list[dict[str, str]]:
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            with path.open(encoding=encoding, newline="") as handle:
                return list(csv.DictReader(handle))
        except UnicodeDecodeError:
            continue
    raise ValueError(f"无法识别 CSV 编码：{path}")


def read_xlsx(path: Path) -> list[dict[str, object]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    rows = list(sheet.iter_rows(values_only=True))
    workbook.close()
    if not rows:
        return []
    headers = [str(value or "").strip() for value in rows[0]]
    return [dict(zip(headers, row)) for row in rows[1:] if any(value is not None for value in row)]


def article_key(url: str) -> str:
    query = parse_qs(urlparse(url).query)
    mid = (query.get("mid") or [""])[0]
    idx = (query.get("idx") or [""])[0]
    sn = (query.get("sn") or [""])[0]
    if mid and idx:
        return f"{mid}:{idx}"
    if sn:
        return sn
    return hashlib.sha256(url.encode("utf-8")).hexdigest()[:24]


def write_csv(path: Path, rows: list[dict[str, object]], fields: list[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields)
        writer.writeheader()
        writer.writerows(rows)


def create_schema(conn: sqlite3.Connection) -> None:
    conn.executescript(
        """
        CREATE TABLE IF NOT EXISTS supplemental_imports (
            batch_id TEXT PRIMARY KEY,
            start_date TEXT NOT NULL,
            end_date TEXT NOT NULL,
            imported_at TEXT NOT NULL,
            raw_dir TEXT NOT NULL,
            export_dir TEXT NOT NULL,
            source_manifest_json TEXT NOT NULL
        );

        CREATE TABLE IF NOT EXISTS search_daily_conversion (
            batch_id TEXT NOT NULL,
            date TEXT NOT NULL,
            converted_followers INTEGER,
            read_users INTEGER,
            raw_json TEXT NOT NULL,
            PRIMARY KEY (batch_id, date)
        );

        CREATE TABLE IF NOT EXISTS search_follower_sources (
            batch_id TEXT NOT NULL,
            search_section TEXT NOT NULL,
            impressions INTEGER,
            clicks INTEGER,
            click_rate REAL,
            converted_followers INTEGER,
            raw_json TEXT NOT NULL,
            PRIMARY KEY (batch_id, search_section)
        );

        CREATE TABLE IF NOT EXISTS search_keywords (
            batch_id TEXT NOT NULL,
            row_no INTEGER NOT NULL,
            keyword TEXT NOT NULL,
            search_section TEXT,
            impressions INTEGER,
            clicks INTEGER,
            click_rate REAL,
            related_terms TEXT,
            raw_json TEXT NOT NULL,
            PRIMARY KEY (batch_id, row_no)
        );

        CREATE TABLE IF NOT EXISTS search_articles (
            batch_id TEXT NOT NULL,
            row_no INTEGER NOT NULL,
            article_key TEXT NOT NULL,
            url TEXT NOT NULL,
            title TEXT,
            impressions INTEGER,
            clicks INTEGER,
            click_rate REAL,
            rank_position REAL,
            search_terms TEXT,
            raw_json TEXT NOT NULL,
            PRIMARY KEY (batch_id, row_no)
        );

        CREATE TABLE IF NOT EXISTS search_service_menus (
            batch_id TEXT NOT NULL,
            row_no INTEGER NOT NULL,
            name TEXT NOT NULL,
            impressions INTEGER,
            clicks INTEGER,
            impression_users INTEGER,
            click_users INTEGER,
            click_rate REAL,
            raw_json TEXT NOT NULL,
            PRIMARY KEY (batch_id, row_no)
        );

        CREATE INDEX IF NOT EXISTS idx_search_daily_date ON search_daily_conversion(date);
        CREATE INDEX IF NOT EXISTS idx_search_keywords_keyword ON search_keywords(keyword);
        CREATE INDEX IF NOT EXISTS idx_search_articles_key ON search_articles(article_key);
        """
    )


def main() -> None:
    parser = argparse.ArgumentParser(description="导入微信公众号搜索与粉丝补充数据")
    parser.add_argument("--batch-id", required=True, help="对应阅读趋势批次 ID")
    parser.add_argument("--follower-sources", required=True, type=Path)
    parser.add_argument("--key-data", required=True, type=Path)
    parser.add_argument("--service-menus", required=True, type=Path)
    parser.add_argument("--keywords", required=True, type=Path)
    parser.add_argument("--articles", required=True, type=Path)
    parser.add_argument("--db-path", type=Path, default=DEFAULT_DB_PATH)
    args = parser.parse_args()

    source_paths = [
        args.follower_sources,
        args.key_data,
        args.service_menus,
        args.keywords,
        args.articles,
    ]
    for path in source_paths:
        if not path.exists():
            raise FileNotFoundError(path)

    conn = sqlite3.connect(str(args.db_path))
    conn.row_factory = sqlite3.Row
    batch = conn.execute(
        "SELECT batch_id FROM import_batches WHERE batch_id = ?", (args.batch_id,)
    ).fetchone()
    if not batch:
        raise ValueError(f"阅读趋势批次不存在：{args.batch_id}")
    dates = re.findall(r"\d{4}-\d{2}-\d{2}", args.key_data.name)
    if len(dates) < 2:
        raise ValueError("无法从关键数据文件名识别起止日期")
    start_date, end_date = dates[0], dates[1]

    raw_dir = DATA_ROOT / "raw" / args.batch_id / "supplemental"
    export_dir = DATA_ROOT / "exports" / args.batch_id / "supplemental"
    raw_dir.mkdir(parents=True, exist_ok=True)
    export_dir.mkdir(parents=True, exist_ok=True)
    manifest = []
    for path in source_paths:
        destination = raw_dir / path.name
        shutil.copy2(path, destination)
        manifest.append(
            {
                "filename": path.name,
                "sha256": sha256_file(path),
                "archived_path": str(destination.relative_to(ROOT)),
            }
        )

    daily_rows = [
        {
            "date": row.get("日期", "").strip(),
            "converted_followers": parse_int(row.get("当日粉丝数")),
            "read_users": parse_int(row.get("当日阅读数")),
        }
        for row in read_csv(args.key_data)
        if row.get("日期")
    ]
    follower_rows = []
    for row in read_xlsx(args.follower_sources):
        section = str(row.get("搜索版块") or "").strip()
        if not section or section == "总计":
            continue
        follower_rows.append(
            {
                "search_section": section,
                "impressions": parse_int(row.get("展示次数")),
                "clicks": parse_int(row.get("点击次数")),
                "click_rate": parse_float(row.get("曝光点击率")),
                "converted_followers": parse_int(row.get("转化粉丝数")),
            }
        )

    keyword_rows = []
    for row in read_csv(args.keywords):
        keyword_rows.append(
            {
                "keyword": (row.get("热门搜索词") or "").strip(),
                "search_section": (row.get("搜索板块") or "").strip(),
                "impressions": parse_int(row.get("展示次数")),
                "clicks": parse_int(row.get("点击次数")),
                "click_rate": parse_float(row.get("曝光点击率")),
                "related_terms": (row.get("关联搜索词") or "").strip(),
            }
        )

    article_rows = []
    for row in read_csv(args.articles):
        url = (row.get("链接") or "").strip()
        rate_value = row.get("曝光有点率") or row.get("曝光点击率")
        article_rows.append(
            {
                "article_key": article_key(url),
                "url": url,
                "title": (row.get("文章名") or "").strip(),
                "impressions": parse_int(row.get("展示次数")),
                "clicks": parse_int(row.get("点击次数")),
                "click_rate": parse_float(rate_value),
                "rank_position": parse_float(row.get("排序位置")),
                "search_terms": (row.get("文章搜索词") or "").strip(),
            }
        )

    menu_rows = []
    for row in read_csv(args.service_menus):
        menu_rows.append(
            {
                "name": (row.get("名称") or "").strip(),
                "impressions": parse_int(row.get("曝光次数")),
                "clicks": parse_int(row.get("点击次数")),
                "impression_users": parse_int(row.get("曝光人数")),
                "click_users": parse_int(row.get("点击人数")),
                "click_rate": parse_float(row.get("曝光点击率")),
            }
        )

    create_schema(conn)
    with conn:
        for table in (
            "search_daily_conversion",
            "search_follower_sources",
            "search_keywords",
            "search_articles",
            "search_service_menus",
        ):
            conn.execute(f"DELETE FROM {table} WHERE batch_id = ?", (args.batch_id,))
        conn.execute("DELETE FROM supplemental_imports WHERE batch_id = ?", (args.batch_id,))
        conn.execute(
            "INSERT INTO supplemental_imports VALUES (?, ?, ?, ?, ?, ?, ?)",
            (
                args.batch_id,
                start_date,
                end_date,
                now(),
                str(raw_dir.relative_to(ROOT)),
                str(export_dir.relative_to(ROOT)),
                json.dumps(manifest, ensure_ascii=False),
            ),
        )
        conn.executemany(
            "INSERT INTO search_daily_conversion VALUES (?, ?, ?, ?, ?)",
            [
                (
                    args.batch_id,
                    row["date"],
                    row["converted_followers"],
                    row["read_users"],
                    json.dumps(row, ensure_ascii=False),
                )
                for row in daily_rows
            ],
        )
        conn.executemany(
            "INSERT INTO search_follower_sources VALUES (?, ?, ?, ?, ?, ?, ?)",
            [
                (
                    args.batch_id,
                    row["search_section"],
                    row["impressions"],
                    row["clicks"],
                    row["click_rate"],
                    row["converted_followers"],
                    json.dumps(row, ensure_ascii=False),
                )
                for row in follower_rows
            ],
        )
        conn.executemany(
            "INSERT INTO search_keywords VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)",
            [
                (
                    args.batch_id,
                    index,
                    row["keyword"],
                    row["search_section"],
                    row["impressions"],
                    row["clicks"],
                    row["click_rate"],
                    row["related_terms"],
                    json.dumps(row, ensure_ascii=False),
                )
                for index, row in enumerate(keyword_rows, start=1)
            ],
        )
        conn.executemany(
            "INSERT INTO search_articles VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)",
            [
                (
                    args.batch_id,
                    index,
                    row["article_key"],
                    row["url"],
                    row["title"],
                    row["impressions"],
                    row["clicks"],
                    row["click_rate"],
                    row["rank_position"],
                    row["search_terms"],
                    json.dumps(row, ensure_ascii=False),
                )
                for index, row in enumerate(article_rows, start=1)
            ],
        )
        conn.executemany(
            "INSERT INTO search_service_menus VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)",
            [
                (
                    args.batch_id,
                    index,
                    row["name"],
                    row["impressions"],
                    row["clicks"],
                    row["impression_users"],
                    row["click_users"],
                    row["click_rate"],
                    json.dumps(row, ensure_ascii=False),
                )
                for index, row in enumerate(menu_rows, start=1)
            ],
        )
    conn.close()

    write_csv(export_dir / "search_daily_conversion.csv", daily_rows, list(daily_rows[0]))
    write_csv(export_dir / "search_follower_sources.csv", follower_rows, list(follower_rows[0]))
    write_csv(export_dir / "search_keywords.csv", keyword_rows, list(keyword_rows[0]))
    write_csv(export_dir / "search_articles.csv", article_rows, list(article_rows[0]))
    write_csv(export_dir / "search_service_menus.csv", menu_rows, list(menu_rows[0]))
    (export_dir / "manifest.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )

    analyzer = ROOT / "scripts" / "analyze_wechat_official_account_stats.py"
    subprocess.run([sys.executable, str(analyzer), "--db-path", str(args.db_path)], check=True)
    print(f"已导入补充数据：{args.batch_id}")
    print(f"原始补充表：{raw_dir.relative_to(ROOT)}")
    print(f"标准化明细：{export_dir.relative_to(ROOT)}")
    print(f"粉丝日数据：{len(daily_rows)} 行；搜索词：{len(keyword_rows)} 行；热门文章：{len(article_rows)} 行")


if __name__ == "__main__":
    main()
